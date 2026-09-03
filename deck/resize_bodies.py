#!/usr/bin/env python3
"""Enlarge the template's body text boxes to use the real slide area.

The template ships body placeholders sized for one line of instruction text.
Filling them with actual content overflows every one. This grows each body box
to the usable band between the title and the footer, splitting it evenly when a
slide has two body boxes, so the content sits inside its shape.
"""
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE

TOP = 1.15        # just under the title bar
BOTTOM = 5.00     # just above the footer rule
LEFT = 0.60
WIDTH = 8.85
GAP = 0.12

prs = Presentation("working.pptx")

# The template carries two kinds of leftover decoration that have to go before
# the real content lands:
#   - empty body text boxes, which render as stray bullet dots;
#   - small empty circles (0.21" square) used as bullet glyphs, one per line of
#     the template's placeholder text. My bullet counts differ, so they end up
#     misaligned and miscounted.
# The APPENDIX divider is also an empty shape, but it is a 10"-wide solid bar,
# so width is what separates them.
for slide in prs.slides:
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            continue                       # carries real text - keep
        w = sh.width / 914400 if sh.width else 0
        is_textbox = sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        # The APPENDIX divider is an empty AUTO_SHAPE too, but it is a 10"-wide
        # solid bar. Only empty text boxes and the small glyphs get removed.
        if is_textbox or w < 0.5:
            sh._element.getparent().remove(sh._element)

# Title slide: the template's three boxes overlap each other out of the box.
title_slide = prs.slides[0]
tshapes = sorted([s for s in title_slide.shapes if s.has_text_frame],
                 key=lambda s: s.top)
if len(tshapes) == 3:
    for sh, (top, hgt) in zip(tshapes, [(1.30, 1.05), (2.45, 0.85), (3.45, 0.40)]):
        sh.left = Inches(1.10)
        sh.top = Inches(top)
        sh.height = Inches(hgt)
        sh.width = Inches(7.90)

for idx, slide in enumerate(prs.slides):
    if idx == 0:
        continue          # the title slide was positioned above; leave it alone
    # Only shapes that actually carry text. These template slides also hold
    # several EMPTY decorative boxes; including them split the band five ways
    # and left every real body 0.67" tall.
    shapes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    if len(shapes) < 2:
        continue
    shapes.sort(key=lambda s: (round(s.top / 914400, 2), round(s.left / 914400, 2)))
    title, rest = shapes[0], shapes[1:]
    # The "Note: you can use more than one slide" boxes sit near the bottom and
    # get cleared anyway - leave them where they are.
    bodies = [s for s in rest if s.top / 914400 < 4.5]
    if not bodies:
        continue
    n = len(bodies)
    band = (BOTTOM - TOP - GAP * (n - 1)) / n
    for i, s in enumerate(bodies):
        s.left = Inches(LEFT)
        s.width = Inches(WIDTH)
        s.top = Inches(TOP + i * (band + GAP))
        s.height = Inches(band)
        s.text_frame.word_wrap = True

prs.save("working.pptx")
print("resized body shapes")
