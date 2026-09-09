#!/usr/bin/env python3
"""Drop the lab screenshots onto their slides, replacing the red placeholders.

Each target slide currently holds one body text box whose first paragraph is a
red "[ SCREENSHOT: ... ]" line followed by an italic note. This removes those
two lines, shrinks the body to the lower part of the slide, and puts the image
in the space that frees up.

Run AFTER replace.py, because replace.py rewrites the text of every shape and
would otherwise discard the pictures.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt

TOP = 1.15        # top of the usable band, matching resize_bodies.py
BOTTOM = 4.82     # stop clear of the footer rule; 5.00 let text touch it
LEFT = 0.60
WIDTH = 8.85
GAP = 0.10
# Sized against the WORST slide (Evaluation for Response, ~158 words of
# observations). Anything taller pushed the last bullet into the footer.
IMG_H = 2.25
BODY_PT = 8.5

# deck slide index -> screenshot for that stage
SHOTS = {
    8:  "screenshots/10_summarization.png",
    10: "screenshots/11_eval_summarization.png",
    13: "screenshots/12_response_generation.png",
    15: "screenshots/13_eval_response.png",
    16: "screenshots/14_extract_outputs.png",
}


def body_with_placeholder(slide):
    """The text box carrying the red screenshot placeholder, if any."""
    for sh in slide.shapes:
        if sh.has_text_frame and "[ SCREENSHOT:" in sh.text_frame.text:
            return sh
    return None


def main(path):
    prs = Presentation(path)
    placed, skipped = [], []

    for idx, img in SHOTS.items():
        slide = prs.slides[idx]
        body = body_with_placeholder(slide)
        if body is None:
            skipped.append((idx, "no placeholder found"))
            continue

        # Drop the placeholder line and the italic note that follows it. Both
        # describe a screenshot that is about to become an actual screenshot.
        #
        # Identify them by their XML element, NOT by the paragraph wrapper:
        # python-pptx builds a fresh _Paragraph object on every access to
        # .paragraphs, so `p not in keep` is true for every paragraph and
        # silently deletes the entire body. That is exactly what it did.
        tf = body.text_frame
        drop = [p._p for p in tf.paragraphs
                if "[ SCREENSHOT:" in p.text
                or p.text.startswith(("Both panels", "Input showing", "Must show"))]
        for el in drop:
            el.getparent().remove(el)
        if not tf.paragraphs or not any(p.text.strip() for p in tf.paragraphs):
            raise SystemExit(f"slide {idx}: removal emptied the body - aborting")

        # Shrink what is left so the observations sit under the image.
        text_top = TOP + IMG_H + GAP
        body.left, body.width = Inches(LEFT), Inches(WIDTH)
        body.top, body.height = Inches(text_top), Inches(BOTTOM - text_top)
        for p in tf.paragraphs:
            p.line_spacing = Pt(11.0)
            for r in p.runs:
                if r.font.size is None or r.font.size > Pt(BODY_PT):
                    r.font.size = Pt(BODY_PT)

        pic = slide.shapes.add_picture(img, Inches(LEFT), Inches(TOP), height=Inches(IMG_H))
        # Centre it: these captures are wider than they are tall, so height is
        # the binding dimension and the width lands wherever it lands.
        pic.left = Inches(LEFT + (WIDTH - pic.width / 914400) / 2)
        placed.append((idx, img, round(pic.width / 914400, 2)))

    prs.save(path)
    for idx, img, w in placed:
        print(f"  slide {idx:2d}  {img.split('/')[-1]:28s} {w}in wide")
    for idx, why in skipped:
        print(f"  slide {idx:2d}  SKIPPED - {why}")
    print(f"{len(placed)} placed, {len(skipped)} skipped")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "ShopNest_Solution_Presentation.pptx")
